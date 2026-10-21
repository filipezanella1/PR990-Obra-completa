"""Exporta a base do app: valores das abas de origem e fórmulas do Painel (arquivo já recalculado)."""
import openpyxl,json,re,sys,os
src=sys.argv[1] if len(sys.argv)>1 else 'volumes_pavimentacao_PR-990.xlsx'
out=sys.argv[2] if len(sys.argv)>2 else 'app_build'
wf=openpyxl.load_workbook(src); wv=openpyxl.load_workbook(src,data_only=True)
P=wf['Painel']
cells={}; refs=set()
for row in P.iter_rows():
    for c in row:
        v=c.value
        if v is None: continue
        if isinstance(v,str) and v.startswith('='):
            cells[c.coordinate]=v
            for s in re.findall(r"'([^']+)'!",v): refs.add(s)
            for s in re.findall(r"(?<![A-Za-z'])([A-Za-zÀ-ú]+)!",v): refs.add(s)
        elif isinstance(v,(int,float,str)): cells[c.coordinate]=v
refs.discard('Painel')
sheets={}
def clean(v):
    if isinstance(v,float): return round(v,10)
    if isinstance(v,(int,str)) or v is None: return v
    return str(v)
for s in sorted(refs):
    ws=wv[s]; data=[]
    for row in ws.iter_rows(min_row=1,max_row=ws.max_row,max_col=ws.max_column,values_only=True):
        data.append([clean(v) for v in row])
    while data and all(v is None for v in data[-1]): data.pop()
    sheets[s]=data
# demais abas de origem (base completa para consulta), só valores
for s in wv.sheetnames:
    if s in ('Painel','Contrato × Projeto','Discrepâncias','Leia-me') or s in sheets: continue
    ws=wv[s]; data=[[clean(v) for v in row] for row in ws.iter_rows(values_only=True)]
    while data and all(v is None for v in data[-1]): data.pop()
    sheets[s]=data
os.makedirs(out,exist_ok=True)
json.dump(dict(cells=cells),open(f'{out}/painel.json','w'),ensure_ascii=False,separators=(',',':'))
json.dump(dict(sheets=sheets,ordem=[s for s in wv.sheetnames if s in sheets],refs=sorted(refs)),open(f'{out}/base.json','w'),ensure_ascii=False,separators=(',',':'))
# valores de referência (para testes): Painel recalculado
ref={c.coordinate:clean(c.value) for row in wv['Painel'].iter_rows() for c in row if c.value is not None}
json.dump(ref,open(f'{out}/ref_painel.json','w'),ensure_ascii=False,separators=(',',':'))
print('células do painel',len(cells),'abas',len(sheets),'refs',sorted(refs))
